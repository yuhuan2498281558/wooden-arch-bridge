from __future__ import annotations

import asyncio
import io
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from bridge_algorithm_service.design_parameter_model import DesignParameterModelUnavailable
from bridge_algorithm_service.main import AnalyzeRequest, ChatRequest, PredictRequest, VisualizeRequest, _draw_bridge, _predict, analyze, annotation_tool, chat, export_excel, health, predict
from bridge_algorithm_service.node_geometry import geometry_payload
from bridge_algorithm_service.node_model import NodeRatioDecision


class MainIntegrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.params = {
            "total_length": 18,
            "span": 15,
            "width": 4.5,
            "n1": 9,
            "n2": 8,
            "rise": 2.7,
        }

    def test_prediction_exposes_backward_compatible_geometry_payload(self) -> None:
        result = _predict(self.params)
        bullhead = result["geometry"]["five_miao_bullhead"]

        self.assertEqual(bullhead["ratios"]["alpha_outer"], 0.666667)
        self.assertEqual(bullhead["ratios"]["beta_inner"], 0.25)
        self.assertEqual(bullhead["source"], "rule_fallback")
        self.assertEqual(bullhead["thickness_mode"], "ignored")
        self.assertIn("SSA-XGBoost + CF-BPNN", result["validation"]["method"])

    def test_prediction_combines_model_parameters_nodes_and_deterministic_lengths(self) -> None:
        result = _predict({**self.params, "rise": None, "total_length": 32})

        self.assertEqual(result["rise_span"], 0.1703)
        self.assertEqual(result["s1_flat_lo"], 251.3)
        self.assertEqual(result["s1_flat_len"], 5.0)
        self.assertAlmostEqual(result["s1_diag_len"], 5.615, places=3)
        self.assertAlmostEqual(result["s2_diag1_len"], 3.743, places=3)
        self.assertAlmostEqual(result["s2_diag2_len"], 3.038, places=3)
        self.assertAlmostEqual(result["s2_flat_len"], 2.5, places=3)
        self.assertEqual(result["validation"]["model_chain"][0]["source"], "mentor_model")
        self.assertEqual(result["validation"]["model_chain"][2]["source"], "deterministic-3plus5-geometry-v1")

    def test_prediction_marks_rule_fallback_when_mentor_model_is_unavailable(self) -> None:
        with patch(
            "bridge_algorithm_service.main.predict_design_parameters",
            side_effect=DesignParameterModelUnavailable("test model missing"),
        ):
            result = _predict(self.params)

        self.assertEqual(result["validation"]["model_status"], "fallback")
        self.assertEqual(result["validation"]["fallback_reason"], "test model missing")
        self.assertEqual(result["validation"]["model_chain"][0]["source"], "rule_fallback")

    def test_prediction_can_expose_pilot_model_ratios_and_metadata(self) -> None:
        decision = NodeRatioDecision(
            alpha=0.61,
            beta=0.22,
            source="pilot_model",
            metadata={
                "model_status": "pilot_not_for_production",
                "model_versions": ["five-miao-node-pilot-v5-test"],
                "within_training_range": True,
            },
        )

        with patch("bridge_algorithm_service.main.resolve_node_ratios", return_value=decision):
            result = _predict(self.params)

        bullhead = result["geometry"]["five_miao_bullhead"]
        self.assertEqual(bullhead["ratios"], {"alpha_outer": 0.61, "beta_inner": 0.22})
        self.assertEqual(bullhead["source"], "pilot_model")
        self.assertEqual(bullhead["model_status"], "pilot_not_for_production")
        self.assertIn("pilot node-ratio model", result["validation"]["method"])

    def test_prediction_passes_designer_mode_and_exposes_mode_source(self) -> None:
        decision = NodeRatioDecision(
            alpha=0.42,
            beta=0.22,
            source="designer_mode_pilot",
            metadata={
                "model_status": "research_designer_mode_not_for_deployment",
                "model_versions": ["five-miao-node-pilot-v9-designer-selected-mode"],
                "design_mode": {
                    "requested": "front_half",
                    "selection_source": "designer_selected",
                    "applied": True,
                },
            },
        )

        with patch("bridge_algorithm_service.main.resolve_node_ratios", return_value=decision) as resolver:
            result = _predict({**self.params, "outer_node_structure_type": "front_half"})

        self.assertEqual(resolver.call_args.kwargs["design_mode"], "front_half")
        bullhead = result["geometry"]["five_miao_bullhead"]
        self.assertEqual(bullhead["source"], "designer_mode_pilot")
        self.assertEqual(bullhead["design_mode"]["selection_source"], "designer_selected")
        self.assertIn("designer-selected mode alpha expert", result["validation"]["method"])

    def test_drawing_accepts_custom_ratios_and_legacy_results(self) -> None:
        legacy_result = _predict(self.params)
        legacy_result.pop("geometry")
        default_png = _draw_bridge(self.params, legacy_result)

        custom_result = _predict(self.params)
        custom_result["geometry"] = geometry_payload(0.58, 0.2, source="test")
        custom_png = _draw_bridge(self.params, custom_result)

        self.assertTrue(default_png.startswith(b"\x89PNG\r\n\x1a\n"))
        self.assertTrue(custom_png.startswith(b"\x89PNG\r\n\x1a\n"))
        self.assertNotEqual(default_png, custom_png)

    def test_annotation_route_points_to_packaged_html(self) -> None:
        response = annotation_tool()
        path = Path(response.path)

        self.assertTrue(path.is_file())
        self.assertEqual(path.name, "index.html")

        html = path.read_text(encoding="utf-8")
        self.assertIn('id="panButton"', html)
        self.assertIn('id="labelButton"', html)
        self.assertIn('id="reviewQueueInput"', html)
        self.assertIn('id="nextReviewButton"', html)
        self.assertIn('id="imageFolderButton"', html)
        self.assertIn('id="imageFolderSelect"', html)
        self.assertIn("right.lastModified - left.lastModified", html)
        self.assertIn("Ctrl+滚轮缩放", html)
        self.assertIn("JSON 中的八点已保留", html)
        self.assertIn("loadedImageFileName === imageFileName", html)

    def test_prediction_endpoint_preserves_multi_span_count(self) -> None:
        response = predict(PredictRequest(
            length=54,
            span=18,
            spans_count=3,
            width=4.5,
            n1=9,
            n2=8,
        ))

        self.assertEqual(response["params"]["spans_count"], 3)
        self.assertEqual(response["params"]["total_length"], 54)
        self.assertEqual(response["params"]["span_arrangement"], "equal-span-standard-units")

    def test_excel_export_contains_named_five_miao_diagonal_lengths(self) -> None:
        result = _predict(self.params)
        response = export_excel(VisualizeRequest(params=self.params, result=result))

        async def collect() -> bytes:
            return b"".join([chunk async for chunk in response.body_iterator])

        body = asyncio.run(collect())
        workbook = load_workbook(io.BytesIO(body), read_only=True)
        sheet = workbook.active
        labels = [str(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True) if row and row[0]]

        self.assertIn("五节苗外斜弦（斜弦1）根径", labels)
        self.assertIn("五节苗内斜弦（斜弦2）根径", labels)
        self.assertIn("五节苗外斜弦（斜弦1）长度", labels)
        self.assertIn("五节苗内斜弦（斜弦2）长度", labels)
        self.assertNotIn("斜弦命名说明", labels)
        self.assertEqual(sheet["A1"].alignment.horizontal, "center")
        self.assertEqual(sheet["A2"].alignment.horizontal, "center")
        workbook.close()


    def test_health_exposes_both_model_stages(self) -> None:
        response = health()

        self.assertTrue(response["design_parameter_model"]["available"])
        self.assertIn("node_model", response)

    def test_chat_uses_rule_router_to_create_design_parameters(self) -> None:
        response = chat(ChatRequest(
            message="设计一座总长54米、净跨18米、桥面宽4.8米、3孔、外节点采用前半区的木拱廊桥",
        ))

        self.assertEqual(response["status"], "ready")
        self.assertEqual(response["params"]["total_length"], 54)
        self.assertEqual(response["params"]["span"], 18)
        self.assertEqual(response["params"]["width"], 4.8)
        self.assertEqual(response["params"]["spans_count"], 3)
        self.assertEqual(response["params"]["outer_node_structure_type"], "front_half")
        self.assertEqual(response["drawing_status"], "blocked")
        self.assertFalse(response["image"])
        self.assertIn("不满足所选结构模式 front_half", response["drawing_error"])
        self.assertIn("node_model_disabled", response["drawing_error"])

    def test_chat_defaults_structure_type_to_back_half(self) -> None:
        response = chat(ChatRequest(
            message="设计一座净跨18米、桥面宽4.8米的木拱廊桥",
        ))

        self.assertEqual(response["status"], "ready")
        self.assertEqual(response["params"]["outer_node_structure_type"], "back_half")
        self.assertEqual(response["drawing_status"], "ready")
        self.assertTrue(response["image"])

    def test_chat_extracts_rise_from_natural_language(self) -> None:
        response = chat(ChatRequest(
            message="设计一座净跨18米、桥面宽4.8米、矢高2.6米的木拱廊桥",
        ))

        self.assertEqual(response["status"], "ready")
        self.assertEqual(response["params"]["rise"], 2.6)
        self.assertEqual(response["drawing_status"], "ready")
        self.assertTrue(response["image"])

    def test_chat_extracts_system_counts_for_n1_and_n2(self) -> None:
        response = chat(ChatRequest(
            message="设计一座净跨18米、桥面宽4.8米、三节苗 8 系、五节苗 7 系的木拱廊桥",
        ))

        self.assertEqual(response["status"], "ready")
        self.assertEqual(response["params"]["n1"], 8)
        self.assertEqual(response["params"]["n2"], 7)
        self.assertEqual(response["drawing_status"], "ready")
        self.assertTrue(response["image"])



    def test_analysis_links_structural_check_and_service_life(self) -> None:
        response = analyze(AnalyzeRequest(
            params=self.params,
            result=_predict(self.params),
            wood_type="杉木",
            service_life_evidence={
                "current_age": 18,
                "exposure_class": "exposed",
                "maintenance_level": "regular",
                "preservative_treatment": False,
            },
        ))["result"]

        self.assertEqual(response["service_life"]["model_version"], "service-life-weibull-gamma-v1")
        self.assertEqual(response["service_life"]["update_mode"], "right_censored_update")
        self.assertGreater(response["service_life"]["estimated_life"], 18)
        self.assertTrue(response["structural_check"]["passes"])
        self.assertIn("selected_decay_life", response["structural_check"])
        self.assertGreater(response["integrated_assessment"]["governing_life_years"], 0)
        self.assertEqual(len(response["integrated_assessment"]["chain"]), 4)
        self.assertEqual(response["member_failure"]["system_dcr"], response["structural_check"]["max_dcr"])

    def test_draw_bridge_generates_valid_png(self) -> None:
        result = _predict(self.params)
        png_bytes = _draw_bridge(self.params, result)
        self.assertIsInstance(png_bytes, bytes)
        self.assertGreater(len(png_bytes), 1000)
        self.assertTrue(png_bytes.startswith(b"\x89PNG\r\n\x1a\n"))

    def test_prediction_binds_normalized_design_input_snapshot(self) -> None:
        result = _predict(self.params)

        self.assertEqual(result["design_inputs"]["span"], 15)
        self.assertEqual(result["design_inputs"]["width"], 4.5)
        self.assertEqual(result["design_inputs"]["spans_count"], 1)

    def test_draw_bridge_rejects_new_params_with_old_result(self) -> None:
        result = _predict(self.params)

        with self.assertRaisesRegex(ValueError, "do not match prediction snapshot: width"):
            _draw_bridge({**self.params, "width": 5.2}, result)

    def test_draw_bridge_rejects_missing_required_specification(self) -> None:
        result = _predict(self.params)
        result.pop("s2_flat_lo")

        with self.assertRaisesRegex(ValueError, "missing numeric field: s2_flat_lo"):
            _draw_bridge(self.params, result)

    def test_draw_bridge_rejects_alpha_outside_selected_mode(self) -> None:
        decision = NodeRatioDecision(
            alpha=0.42,
            beta=0.22,
            source="designer_mode_pilot",
            metadata={
                "model_status": "research_designer_mode_not_for_deployment",
                "model_versions": ["five-miao-node-pilot-v9-designer-selected-mode"],
                "within_training_range": True,
                "within_local_domain": True,
                "target_models": {},
                "model_warnings": [],
                "design_mode": {
                    "requested": "front_half",
                    "selection_source": "designer_selected",
                    "applied": True,
                },
            },
        )
        params = {**self.params, "outer_node_structure_type": "front_half"}
        with patch("bridge_algorithm_service.main.resolve_node_ratios", return_value=decision):
            result = _predict(params)
        result["geometry"]["five_miao_bullhead"]["ratios"]["alpha_outer"] = 0.58

        with self.assertRaisesRegex(ValueError, "不满足所选结构模式 front_half"):
            _draw_bridge(params, result)


if __name__ == "__main__":
    unittest.main()
