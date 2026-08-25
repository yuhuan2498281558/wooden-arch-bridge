from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import build_delivery_report as base


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "delivery" / "wood-arch-bridge-system-report-current.md"
OUTPUT = ROOT / "docs" / "delivery" / "中国木拱廊桥智能设计系统_系统技术与使用说明_V1.1_精排版.docx"
ASSET_DIR = ROOT / "docs" / "delivery" / "_report_assets_polished"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

# standard_business_brief preset + named engineering-delivery overrides.
NAVY = "102A43"
BLUE = "1F5F99"
SKY = "DDEAF4"
PALE = "F3F7FA"
GOLD = "C8923E"
INK = "243B53"
MUTED = "627D98"
GRID = "C9D5DF"
WHITE = "FFFFFF"
GREEN = "E8F4EC"
GREEN_TEXT = "2B6B45"
AMBER = "FFF4D8"
AMBER_TEXT = "7A5A00"
FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"


def pc(color: str) -> str:
    return color if color.startswith("#") else f"#{color}"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    paths = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]
    for path in paths:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def create_cover(path: Path) -> None:
    w, h = 1950, 2700
    img = Image.new("RGB", (w, h), pc(NAVY))
    draw = ImageDraw.Draw(img)

    # Architectural grid and a quiet diagonal light field.
    for x in range(0, w, 130):
        draw.line((x, 0, x, h), fill="#173A57", width=2)
    for y in range(0, h, 130):
        draw.line((0, y, w, y), fill="#173A57", width=2)
    draw.polygon([(1180, 0), (w, 0), (w, 1460), (780, 2700), (220, 2700)], fill="#14334F")
    draw.rectangle((0, 0, 34, h), fill=pc(GOLD))

    draw.text((150, 170), "SYSTEM DELIVERY REPORT", font=font(42, True), fill="#A9C4D8")
    draw.text((150, 258), "中国木拱廊桥", font=font(92, True), fill=pc(WHITE))
    draw.text((150, 380), "智能设计系统", font=font(92, True), fill=pc(WHITE))
    draw.rectangle((150, 520, 620, 530), fill=pc(GOLD))
    draw.text((150, 585), "系统技术与使用说明", font=font(54, True), fill="#F2D4A3")
    draw.text((150, 670), "架构 · 技术 · 功能 · 操作 · 部署 · 验证", font=font(31), fill="#A9C4D8")

    # Stylized timber arch and structural node system.
    left, right, deck_y = 180, 1770, 1850
    draw.line((left, deck_y, right, deck_y), fill="#EAF2F8", width=14)
    draw.arc((left, 965, right, 2320), 192, 348, fill="#F5F9FC", width=28)
    draw.arc((360, 1130, 1590, 2185), 192, 348, fill="#7FA9C4", width=18)
    nodes = [left, 445, 710, 975, 1240, 1505, right]
    for idx, x in enumerate(nodes):
        frac = abs(x - (left + right) / 2) / ((right - left) / 2)
        top_y = 1850 - int(500 * (1 - frac**1.7))
        draw.line((x, deck_y, x, top_y), fill="#A9C4D8", width=11)
        draw.ellipse((x - 15, top_y - 15, x + 15, top_y + 15), fill=pc(GOLD))
    for i in range(len(nodes) - 1):
        x1, x2 = nodes[i], nodes[i + 1]
        draw.line((x1, deck_y, x2, deck_y - 250 + (i % 2) * 70), fill="#567F9B", width=8)

    draw.text((150, 2190), "现状交付版", font=font(34, True), fill=pc(GOLD))
    draw.text((150, 2270), "VERSION 1.1  ·  2026.08.13", font=font(31, True), fill=pc(WHITE))
    draw.text((150, 2380), "面向建设方、业务使用人员、系统运维及后续研发", font=font(28), fill="#B9CCDA")
    draw.rectangle((150, 2510, 1800, 2514), fill="#345670")
    draw.text((150, 2560), "WOOD ARCH BRIDGE INTELLIGENT DESIGN SYSTEM", font=font(24, True), fill="#7898AE")
    img.save(path, quality=95)


def rounded(draw: ImageDraw.ImageDraw, xy, fill, outline=None, radius=24, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=pc(fill), outline=pc(outline) if outline else None, width=width)


def create_architecture(path: Path) -> None:
    w, h = 1900, 1120
    img = Image.new("RGB", (w, h), "#F7FAFC")
    d = ImageDraw.Draw(img)
    d.text((70, 48), "系统逻辑架构", font=font(48, True), fill=pc(NAVY))
    d.text((70, 112), "前端交互、平台管理与桥梁算法服务解耦", font=font(25), fill=pc(MUTED))

    rounded(d, (100, 205, 1800, 350), WHITE, "#8FB3CA", 30, 3)
    d.rectangle((100, 205, 123, 350), fill=pc(GOLD))
    d.text((160, 244), "浏览器端", font=font(34, True), fill=pc(NAVY))
    d.text((420, 250), "设计工作台 · 动态菜单/RBAC · 图纸/结果 · 文件下载", font=font(28), fill=pc(INK))

    rounded(d, (695, 435, 1205, 550), SKY, "#5B8EB1", 28, 3)
    d.text((820, 470), "Nginx 统一入口", font=font(31, True), fill=pc(NAVY))
    d.line((950, 350, 950, 435), fill="#5B8EB1", width=8)

    rounded(d, (105, 670, 790, 855), WHITE, "#5B8EB1", 28, 3)
    d.text((160, 707), "Django / DRF 平台", font=font(32, True), fill=pc(NAVY))
    d.text((160, 765), "认证 · 用户 · 角色 · 菜单 · 管理", font=font(25), fill=pc(MUTED))
    rounded(d, (1110, 670, 1795, 855), WHITE, "#5B8EB1", 28, 3)
    d.text((1165, 707), "FastAPI 算法服务", font=font(32, True), fill=pc(NAVY))
    d.text((1165, 765), "规则预测 · 绘图 · 分析 · 优化 · 导出", font=font(25), fill=pc(MUTED))
    d.line((735, 550, 450, 670), fill="#5B8EB1", width=7)
    d.line((1165, 550, 1450, 670), fill="#5B8EB1", width=7)
    d.text((520, 568), "/api/", font=font(23), fill=pc(MUTED))
    d.text((1255, 568), "/bridge-api/", font=font(23), fill=pc(MUTED))

    rounded(d, (105, 945, 790, 1055), "EDF3F7", "#B4C5D0", 24, 2)
    d.text((160, 977), "PostgreSQL（外部） · Redis", font=font(27, True), fill=pc(INK))
    rounded(d, (1110, 945, 1795, 1055), "EDF3F7", "#B4C5D0", 24, 2)
    d.text((1165, 970), "BIMFACE（可选）", font=font(27, True), fill=pc(INK))
    d.text((1450, 980), "预配置模型查看", font=font(22), fill=pc(MUTED))
    d.line((450, 855, 450, 945), fill="#8FA4B4", width=5)
    d.line((1450, 855, 1450, 945), fill="#8FA4B4", width=5)
    img.save(path)


def create_workflow(path: Path) -> None:
    w, h = 1900, 600
    img = Image.new("RGB", (w, h), pc(WHITE))
    d = ImageDraw.Draw(img)
    d.text((60, 40), "现有方案工作流", font=font(44, True), fill=pc(NAVY))
    items = [
        ("01", "输入", "自然语言 / 表单"),
        ("02", "归一化", "边界 / 默认值"),
        ("03", "规则估算", "几何 / 根径 / 长度"),
        ("04", "成果生成", "图纸 / Excel"),
        ("05", "辅助分析", "安全 / 寿命 / 候选"),
    ]
    start, top, bw, gap = 60, 165, 310, 60
    for i, (no, title, desc) in enumerate(items):
        x = start + i * (bw + gap)
        rounded(d, (x, top, x + bw, top + 280), PALE, "#A7BECE", 25, 3)
        d.rectangle((x, top, x + 12, top + 280), fill=pc(GOLD))
        d.text((x + 32, top + 32), no, font=font(27, True), fill=pc(GOLD))
        d.text((x + 32, top + 92), title, font=font(31, True), fill=pc(NAVY))
        d.text((x + 32, top + 156), desc, font=font(23), fill=pc(INK))
        if i < len(items) - 1:
            ax = x + bw + 10
            d.line((ax, top + 140, ax + 32, top + 140), fill="#5B8EB1", width=7)
            d.polygon([(ax + 32, top + 127), (ax + 52, top + 140), (ax + 32, top + 153)], fill="#5B8EB1")
    d.text((60, 515), "离线研发：八点标注 → 数据修复 → 对称目标 → 分组验证 → Pilot（未接入生产）", font=font(25), fill=pc(MUTED))
    img.save(path)


def create_modules(path: Path) -> None:
    w, h = 1900, 760
    img = Image.new("RGB", (w, h), "#F7FAFC")
    d = ImageDraw.Draw(img)
    d.text((65, 42), "六大业务功能区", font=font(46, True), fill=pc(NAVY))
    d.text((65, 106), "同一工作台、同一方案状态、路由与页签双向同步", font=font(25), fill=pc(MUTED))
    items = [
        ("01", "3D 模型", "BIMFACE 查看入口"),
        ("02", "结构图纸", "生成 · 缩放 · 下载"),
        ("03", "设计参数", "录入 · 校验 · 重算"),
        ("04", "预测结果", "几何 · 根径 · 长度"),
        ("05", "安全分析", "风险 · 寿命 · 提示"),
        ("06", "优化设计", "候选 · 比较 · 回填"),
    ]
    for i, (no, title, desc) in enumerate(items):
        col, row = i % 3, i // 3
        x = 65 + col * 610
        y = 190 + row * 245
        rounded(d, (x, y, x + 555, y + 190), WHITE, "#B0C4D3", 24, 2)
        d.rectangle((x, y, x + 12, y + 190), fill=pc(GOLD if i in (1, 4) else BLUE))
        d.text((x + 34, y + 28), no, font=font(25, True), fill=pc(GOLD))
        d.text((x + 105, y + 25), title, font=font(31, True), fill=pc(NAVY))
        d.text((x + 34, y + 105), desc, font=font(23), fill=pc(MUTED))
    img.save(path)


def set_run(run, *, size=11, color=INK, bold=None, italic=None, latin=FONT_LATIN, cjk=FONT_CJK):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cjk)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_border(paragraph, *, side: str, color: str, size: int, space: int = 6):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)
    p_bdr.append(edge)


def shade_paragraph(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT_LATIN
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    h1.font.size = Pt(18)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(WHITE)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(14)
    h1.paragraph_format.line_spacing = 1.0
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.left_indent = Inches(0.0)

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT_LATIN
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(BLUE)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True

    h3 = doc.styles["Heading 3"]
    h3.font.name = FONT_LATIN
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string("1F4D78")
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.keep_together = True

    for style_name in ("TOC 1", "TOC 2"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT_LATIN
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
            style.font.size = Pt(11 if style_name == "TOC 1" else 9.5)
            style.font.color.rgb = RGBColor.from_string(NAVY if style_name == "TOC 1" else MUTED)
            style.paragraph_format.space_after = Pt(7 if style_name == "TOC 1" else 3)


def add_field(run, instruction: str, placeholder: str = ""):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def set_last_image_alt(doc: Document, title: str, description: str) -> None:
    inline = doc.inline_shapes[-1]._inline
    inline.docPr.set("title", title)
    inline.docPr.set("descr", description)


def configure_page(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    sec.different_first_page_header_footer = True

    header_p = sec.header.paragraphs[0]
    header_p.paragraph_format.space_after = Pt(4)
    header_p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = header_p.add_run("中国木拱廊桥智能设计系统")
    set_run(r, size=8.5, color=MUTED, bold=True)
    r = header_p.add_run("\tSYSTEM TECHNICAL & USER GUIDE")
    set_run(r, size=7.5, color=MUTED, bold=True)
    set_border(header_p, side="bottom", color="D6E0E8", size=5, space=5)

    footer_p = sec.footer.paragraphs[0]
    footer_p.paragraph_format.space_before = Pt(4)
    footer_p.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    footer_p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r = footer_p.add_run("现状交付版 · V1.1")
    set_run(r, size=8.5, color=MUTED)
    r = footer_p.add_run("\tWOOD ARCH BRIDGE")
    set_run(r, size=7.5, color="8BA1B2", bold=True)
    r = footer_p.add_run("\t")
    set_run(r, size=8.5, color=MUTED)
    add_field(r, " PAGE ", "1")

    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)


def add_cover_page(doc: Document, image: Path) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(image), width=Inches(6.5), height=Inches(8.65))
    set_last_image_alt(doc, "报告封面", "中国木拱廊桥智能设计系统技术与使用说明封面，含木拱结构示意图。")
    doc.add_page_break()


def add_front_title(doc: Document, kicker: str, title: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(kicker.upper())
    set_run(r, size=8.5, color=GOLD, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title)
    set_run(r, size=22, color=NAVY, bold=True)
    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(subtitle)
        set_run(r, size=10.5, color=MUTED)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(14)
    set_border(rule, side="bottom", color=GOLD, size=14, space=0)


def add_inline(paragraph, text: str, size=11, color=INK):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run(run, size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run(run, size=max(9, size - 1), color=BLUE, latin="Consolas")
        else:
            run = paragraph.add_run(part)
            set_run(run, size=size, color=color)


def style_body(paragraph, *, after=6, line=1.10, keep=False):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_together = keep


def add_callout(doc: Document, text: str, label: str = "交付说明") -> None:
    p = doc.add_paragraph()
    style_body(p, after=12, line=1.15, keep=True)
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.08)
    shade_paragraph(p, "EDF4F8")
    set_border(p, side="left", color=GOLD, size=18, space=8)
    r = p.add_run(f"{label}  ")
    set_run(r, size=10, color=NAVY, bold=True)
    add_inline(p, text, size=10, color=INK)


def set_table_borders(table, color=GRID, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)
        borders.append(edge)


def column_widths(rows: list[list[str]], cols: int) -> list[int]:
    if cols == 2:
        return [2200, 7160]
    if cols == 3:
        return [1500, 3100, 4760]
    if cols == 4:
        return [1700, 2750, 1450, 3460]
    base_width = 9360 // cols
    return [base_width] * (cols - 1) + [9360 - base_width * (cols - 1)]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows = [rows[0]] + rows[2:]
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    for i, row in enumerate(rows):
        for j in range(cols):
            text = row[j] if j < len(row) else ""
            if text == "V1.0":
                text = "V1.1（版式优化）"
            cell = table.cell(i, j)
            cell.text = ""
            base.set_cell_margins(cell, top=110, start=140, bottom=110, end=140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            style_body(p, after=0, line=1.10)
            add_inline(p, text, size=9.3, color=INK)
            if i == 0:
                base.set_cell_shading(cell, NAVY)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                if i % 2 == 0:
                    base.set_cell_shading(cell, "F7F9FB")
                if text in {"已具备", "已具备 V1", "工具已具备"}:
                    base.set_cell_shading(cell, GREEN)
                    for run in p.runs:
                        run.font.color.rgb = RGBColor.from_string(GREEN_TEXT)
                        run.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif text == "条件具备":
                    base.set_cell_shading(cell, AMBER)
                    for run in p.runs:
                        run.font.color.rgb = RGBColor.from_string(AMBER_TEXT)
                        run.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif j > 0 and len(text) <= 12 and not any(ch in text for ch in "，。；："):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_repeat_table_header(table.rows[0])
    base.set_table_geometry(table, column_widths(rows, cols))
    for row in table.rows:
        base.set_cant_split(row)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def new_numbering(doc: Document, fmt: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if fmt == "bullet" else "decimal")
    lvl.append(num_fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•" if fmt == "bullet" else "%1.")
    lvl.append(text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    lvl.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "360")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    if fmt == "bullet":
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
        r_pr.append(fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    first_num = next((idx for idx, child in enumerate(numbering) if child.tag == qn("w:num")), len(numbering))
    numbering.insert(first_num, abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(num)
    p_pr.append(num_pr)


def add_heading(doc: Document, text: str, level: int, workflow: Path, modules: Path) -> None:
    p = doc.add_heading(text, level=level)
    if level == 1:
        shade_paragraph(p, NAVY)
        set_border(p, side="left", color=GOLD, size=24, space=10)
        p.paragraph_format.left_indent = Inches(0.08)
        p.paragraph_format.right_indent = Inches(0.0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(14)
        for run in p.runs:
            run.font.color.rgb = RGBColor.from_string(WHITE)
    elif level == 2:
        set_border(p, side="left", color=GOLD, size=12, space=7)
        p.paragraph_format.left_indent = Inches(0.06)
    if text.startswith(("6.10 ", "9.3 ")):
        p.paragraph_format.page_break_before = True
    if text.startswith("2.4 "):
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(10)
        p2.add_run().add_picture(str(workflow), width=Inches(6.38))
        set_last_image_alt(doc, "现有方案工作流", "从参数输入、归一化、规则估算到成果生成和辅助分析的五步流程。")
    if text.startswith("4 "):
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(12)
        p2.add_run().add_picture(str(modules), width=Inches(6.38))
        set_last_image_alt(doc, "六大业务功能区", "3D模型、结构图纸、设计参数、预测结果、安全分析和优化设计六个功能区。")


def add_code_block(doc: Document, code: str, architecture: Path) -> None:
    if "浏览器端" in code and "Nginx" in code:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        p.add_run().add_picture(str(architecture), width=Inches(6.38))
        set_last_image_alt(doc, "系统逻辑架构", "浏览器经Nginx分别访问Django平台和FastAPI算法服务，连接数据基础设施和可选BIMFACE服务。")
        return
    lines = code.rstrip().splitlines()
    for idx, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        shade_paragraph(p, "F1F4F6")
        if idx == 0:
            set_border(p, side="left", color=GOLD, size=14, space=6)
        r = p.add_run(line or " ")
        set_run(r, size=8.6, color="334E68", latin="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_toc(doc: Document) -> None:
    add_front_title(doc, "NAVIGATION", "目录", "按章节快速定位系统架构、功能、操作和交付边界")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run()
    add_field(r, ' TOC \\o "1-1" \\h \\z \\u ', "目录将在 Word 中自动更新")
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(22)
    note.paragraph_format.space_after = Pt(0)
    set_border(note, side="top", color="D6E0E8", size=5, space=8)
    r = note.add_run("提示：本报告仅描述基准日期时已经实现并验证的内容；研究原型与待配置能力均单列边界。")
    set_run(r, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def extract_front_table(lines: list[str]) -> list[list[str]]:
    start = lines.index("## 文档控制") + 1
    table_lines: list[str] = []
    for line in lines[start:]:
        if line.startswith("|"):
            table_lines.append(line)
        elif table_lines:
            break
    return base.parse_table(table_lines)


def parse_body(doc: Document, lines: list[str], architecture: Path, workflow: Path, modules: Path) -> None:
    start = next(i for i, line in enumerate(lines) if line.startswith("# 1 "))
    lines = lines[start:]
    i = 0
    in_code = False
    code_lines: list[str] = []
    current_num: int | None = None
    current_bullet: int | None = None
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            current_num = current_bullet = None
            if in_code:
                add_code_block(doc, "\n".join(code_lines), architecture)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith("|"):
            current_num = current_bullet = None
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, base.parse_table(table_lines))
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            current_num = current_bullet = None
            add_heading(doc, heading.group(2), len(heading.group(1)), workflow, modules)
            i += 1
            continue
        if line.startswith("> "):
            current_num = current_bullet = None
            add_callout(doc, line[2:].strip())
            i += 1
            continue
        if re.match(r"^-\s+", line):
            current_num = None
            if current_bullet is None:
                current_bullet = new_numbering(doc, "bullet")
            p = doc.add_paragraph()
            style_body(p, after=8, line=1.167)
            apply_numbering(p, current_bullet)
            add_inline(p, re.sub(r"^-\s+", "", line))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            current_bullet = None
            if current_num is None:
                current_num = new_numbering(doc, "decimal")
            p = doc.add_paragraph()
            style_body(p, after=8, line=1.167)
            apply_numbering(p, current_num)
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        current_num = current_bullet = None
        p = doc.add_paragraph()
        style_body(p)
        add_inline(p, line)
        i += 1


def build() -> Path:
    cover = ASSET_DIR / "cover-polished.png"
    architecture = ASSET_DIR / "architecture-polished.png"
    workflow = ASSET_DIR / "workflow-polished.png"
    modules = ASSET_DIR / "modules-polished.png"
    create_cover(cover)
    create_architecture(architecture)
    create_workflow(workflow)
    create_modules(modules)

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    props = doc.core_properties
    props.title = "中国木拱廊桥智能设计系统——系统技术与使用说明（精排版）"
    props.subject = "系统现状交付报告"
    props.author = "项目交付组"
    props.keywords = "木拱廊桥, 系统架构, 技术说明, 使用说明, 部署, 交付"
    props.comments = "V1.1版式优化；事实基准日期2026-08-13"

    add_cover_page(doc, cover)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    add_front_title(doc, "DOCUMENT PROFILE", "文档概览", "交付范围、事实边界与阅读说明")
    add_table(doc, extract_front_table(lines))
    add_callout(
        doc,
        "本报告先覆盖系统已经具备的部分。标记为“待确认”“待标定”或“研究原型”的内容，不计入当前正式功能承诺。",
        label="事实边界",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("阅读导览")
    set_run(r, size=13, color=BLUE, bold=True)
    guide_items = [
        "建设方：优先阅读第 1、2、8、9 章，了解交付范围、架构、验证和边界。",
        "业务使用人员：重点阅读第 4、6 章，掌握功能与完整操作流程。",
        "运维人员：重点阅读第 3、7 章，核对技术栈、部署配置和健康检查。",
        "研发人员：重点阅读第 5 章和附录，了解算法、数据链路及接口依据。",
    ]
    bullet_id = new_numbering(doc, "bullet")
    for item in guide_items:
        p = doc.add_paragraph()
        style_body(p, after=8, line=1.167)
        apply_numbering(p, bullet_id)
        add_inline(p, item)
    doc.add_page_break()
    add_toc(doc)
    parse_body(doc, lines, architecture, workflow, modules)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
