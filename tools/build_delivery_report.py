from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "delivery" / "wood-arch-bridge-system-report-current.md"
OUTPUT = ROOT / "docs" / "delivery" / "中国木拱廊桥智能设计系统_系统技术与使用说明_V1.0.docx"
ASSET_DIR = ROOT / "docs" / "delivery" / "_report_assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

BLUE = "1F4E78"
MID_BLUE = "2E74B5"
PALE_BLUE = "E8EEF5"
NAVY = "162B43"
INK = "22313F"
MUTED = "657586"
WARM = "C88A4A"
WHITE = "FFFFFF"
GRID = "B8C5D1"
LIGHT = "F5F7FA"
FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"


def pc(color: str) -> str:
    return color if color.startswith("#") else f"#{color}"


def cjk_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def rounded_box(draw: ImageDraw.ImageDraw, xy, fill, outline=None, radius=24, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=pc(fill), outline=pc(outline) if outline else None, width=width)


def create_cover_art(path: Path) -> None:
    w, h = 1800, 680
    img = Image.new("RGB", (w, h), pc(WHITE))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, w, h), fill=pc(NAVY))
    for x in range(-200, w + 200, 160):
        d.line((x, h, x + 420, 0), fill="#203B58", width=2)
    d.rectangle((0, 0, 20, h), fill=pc(WARM))
    d.arc((170, 155, 1630, 1080), 190, 350, fill="#D7E6F3", width=22)
    d.arc((360, 260, 1440, 930), 190, 350, fill="#86A9C4", width=13)
    d.line((170, 490, 1630, 490), fill="#D7E6F3", width=10)
    for x in (170, 535, 900, 1265, 1630):
        y = 490 - int(165 * (1 - abs(x - 900) / 730))
        d.line((x, 490, x, y), fill="#86A9C4", width=8)
    d.text((125, 70), "SYSTEM DELIVERY REPORT", font=cjk_font(34, True), fill="#86A9C4")
    d.text((125, 120), "木拱廊桥 · 参数化设计 · 工程表达", font=cjk_font(40, True), fill=pc(WHITE))
    img.save(path)


def create_architecture(path: Path) -> None:
    w, h = 1800, 980
    img = Image.new("RGB", (w, h), "#F7F9FC")
    d = ImageDraw.Draw(img)
    title = cjk_font(42, True)
    body = cjk_font(30)
    small = cjk_font(24)
    d.text((70, 45), "系统逻辑架构", font=title, fill="#183B56")

    rounded_box(d, (95, 150, 1705, 280), WHITE, "#86A9C4", 28, 3)
    d.text((145, 180), "浏览器端", font=cjk_font(31, True), fill="#183B56")
    d.text((365, 180), "设计工作台 · 动态菜单/RBAC · 结果展示 · 图纸与 Excel 下载", font=body, fill=pc(INK))

    rounded_box(d, (650, 340, 1150, 440), "#DDEAF4", "#4C7FA3", 25, 3)
    d.text((815, 368), "Nginx 网关", font=cjk_font(31, True), fill="#183B56")
    d.line((900, 280, 900, 340), fill="#4C7FA3", width=7)

    rounded_box(d, (120, 535, 760, 700), WHITE, "#4C7FA3", 28, 3)
    d.text((175, 568), "Django / DRF", font=cjk_font(31, True), fill="#183B56")
    d.text((175, 615), "认证、用户、角色、菜单、平台管理", font=small, fill=pc(INK))
    rounded_box(d, (1040, 535, 1680, 700), WHITE, "#4C7FA3", 28, 3)
    d.text((1095, 568), "FastAPI 算法服务", font=cjk_font(31, True), fill="#183B56")
    d.text((1095, 615), "规则预测、绘图、分析、优化、导出", font=small, fill=pc(INK))
    d.line((780, 440, 440, 535), fill="#4C7FA3", width=6)
    d.line((1020, 440, 1360, 535), fill="#4C7FA3", width=6)
    d.text((530, 454), "/api/", font=small, fill=pc(MUTED))
    d.text((1180, 454), "/bridge-api/", font=small, fill=pc(MUTED))

    rounded_box(d, (120, 795, 760, 910), "#EDF2F6", "#A9B9C6", 24, 2)
    d.text((175, 828), "PostgreSQL（外部） · Redis", font=cjk_font(27, True), fill="#183B56")
    rounded_box(d, (1040, 795, 1680, 910), "#EDF2F6", "#A9B9C6", 24, 2)
    d.text((1095, 820), "BIMFACE（可选外部服务）", font=cjk_font(27, True), fill="#183B56")
    d.text((1095, 860), "预配置模型查看", font=small, fill=pc(MUTED))
    d.line((440, 700, 440, 795), fill="#8299AA", width=5)
    d.line((1360, 700, 1360, 795), fill="#8299AA", width=5)
    img.save(path)


def create_workflow(path: Path) -> None:
    w, h = 1800, 560
    img = Image.new("RGB", (w, h), pc(WHITE))
    d = ImageDraw.Draw(img)
    d.text((70, 30), "现有方案工作流", font=cjk_font(40, True), fill="#183B56")
    labels = [
        ("01", "输入", "自然语言 / 表单"),
        ("02", "归一化", "边界与默认值"),
        ("03", "规则估算", "几何 / 根径 / 长度"),
        ("04", "成果", "图纸 / Excel"),
        ("05", "辅助分析", "安全 / 寿命 / 候选"),
    ]
    x0, y0, bw, gap = 70, 150, 290, 60
    for i, (num, title, detail) in enumerate(labels):
        x = x0 + i * (bw + gap)
        rounded_box(d, (x, y0, x + bw, y0 + 250), "#F4F7FA", "#98AFC1", 25, 3)
        d.ellipse((x + 22, y0 + 22, x + 82, y0 + 82), fill=pc(WARM))
        d.text((x + 34, y0 + 31), num, font=cjk_font(22, True), fill=pc(WHITE))
        d.text((x + 25, y0 + 110), title, font=cjk_font(30, True), fill="#183B56")
        d.text((x + 25, y0 + 163), detail, font=cjk_font(23), fill=pc(INK))
        if i < len(labels) - 1:
            ax = x + bw + 10
            d.line((ax, y0 + 125, ax + 38, y0 + 125), fill="#4C7FA3", width=6)
            d.polygon([(ax + 38, y0 + 112), (ax + 56, y0 + 125), (ax + 38, y0 + 138)], fill="#4C7FA3")
    d.text((70, 455), "离线研发链路：八点标注 → 元数据修复 → 对称目标 → 分组验证 → Pilot 模型（未接入生产）", font=cjk_font(25), fill=pc(MUTED))
    img.save(path)


def set_run_font(run, name=FONT_LATIN, east_asia=FONT_CJK, size=None, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def new_decimal_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    level.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    level.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "269")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "271")
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    first_num_index = next(
        (index for index, child in enumerate(numbering) if child.tag == qn("w:num")),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(num)
    p_pr.append(num_pr)


def set_table_geometry(table, widths: list[int]):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        set_cant_split(row)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table_widths(cols: int) -> list[int]:
    if cols == 2:
        return [2200, 7160]
    if cols == 3:
        return [2200, 3000, 4160]
    if cols == 4:
        return [1650, 2500, 1550, 3660]
    base = 9360 // cols
    return [base] * (cols - 1) + [9360 - base * (cols - 1)]


def add_inline(paragraph, text: str, size=11, color=INK):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", east_asia="Microsoft YaHei", size=max(9, size - 1), color=BLUE)
            run.font.highlight_color = None
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def style_paragraph(paragraph, *, after=6, line=1.25, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = keep


def add_callout(doc: Document, text: str):
    p = doc.add_paragraph()
    style_paragraph(p, after=10, line=1.25, keep=True)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EEF4F8")
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), MID_BLUE)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    run = p.add_run("说明  ")
    set_run_font(run, size=10.5, color=BLUE, bold=True)
    add_inline(p, text, 10.5, INK)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, MID_BLUE, 18, 10),
        ("Heading 2", 13, MID_BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    doc.styles["Heading 1"].paragraph_format.page_break_before = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.4)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("中国木拱廊桥智能设计系统  ·  系统技术与使用说明")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "C7D3DE")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    add_page_field(section.footer.paragraphs[0])


def add_cover(doc: Document, cover_art: Path):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("中国木拱廊桥智能设计系统")
    set_run_font(r, size=25, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("系统技术与使用说明")
    set_run_font(r, size=18, color=MID_BLUE, bold=True)

    doc.add_picture(str(cover_art), width=Inches(6.5))
    pic_p = doc.paragraphs[-1]
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_after = Pt(24)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.style = "Table Grid"
    values = [
        ("文档性质", "现状交付版"),
        ("版本", "V1.0"),
        ("基准日期", "2026 年 8 月 13 日"),
        ("适用对象", "建设方、业务用户、运维与后续研发人员"),
    ]
    for i, (label, value) in enumerate(values):
        meta.cell(i, 0).text = label
        meta.cell(i, 1).text = value
        set_cell_shading(meta.cell(i, 0), PALE_BLUE)
        for cell in meta.rows[i].cells:
            for p2 in cell.paragraphs:
                style_paragraph(p2, after=0, line=1.1)
                for run in p2.runs:
                    set_run_font(run, size=10, color=INK, bold=(cell is meta.cell(i, 0)))
    set_table_geometry(meta, [1900, 7460])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("交付信息待确认后补充建设单位、项目编号与签批页")
    set_run_font(r, size=9, color=MUTED)
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int, arch_path: Path, flow_path: Path):
    p = doc.add_heading(text, level=level)
    if level == 1:
        p.paragraph_format.page_break_before = True
    if text.strip().startswith(("6.10 ", "9.3 ")):
        p.paragraph_format.page_break_before = True
    if text.strip().startswith("2.4"):
        doc.add_picture(str(flow_path), width=Inches(6.45))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        vals = [v.strip() for v in line.strip().strip("|").split("|")]
        rows.append(vals)
    return rows


def add_table(doc: Document, rows: list[list[str]]):
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows = [rows[0]] + rows[2:]
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            text = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            style_paragraph(p, after=0, line=1.12)
            add_inline(p, text, size=9.2, color=INK)
            if i == 0:
                set_cell_shading(cell, PALE_BLUE)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(BLUE)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, table_widths(cols))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, code: str, architecture_path: Path):
    if "浏览器端" in code and "Nginx" in code:
        doc.add_picture(str(architecture_path), width=Inches(6.45))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.space_after = Pt(8)
        return
    for line in code.rstrip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F3F5F7")
        p_pr.append(shd)
        r = p.add_run(line or " ")
        set_run_font(r, name="Consolas", east_asia="Microsoft YaHei", size=8.5, color="334E68")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build() -> Path:
    cover = ASSET_DIR / "cover-art.png"
    architecture = ASSET_DIR / "architecture.png"
    workflow = ASSET_DIR / "workflow.png"
    create_cover_art(cover)
    create_architecture(architecture)
    create_workflow(workflow)

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    core = doc.core_properties
    core.title = "中国木拱廊桥智能设计系统——系统技术与使用说明"
    core.subject = "现状交付版系统报告"
    core.author = "项目交付组"
    core.keywords = "木拱廊桥, 系统架构, 技术说明, 使用说明, 交付"
    core.comments = "以 2026-08-13 当前代码和验证状态为事实边界"
    add_cover(doc, cover)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = lines.index("## 文档控制")
    lines = lines[start:]
    i = 0
    in_code = False
    code_lines: list[str] = []
    current_num_id: int | None = None
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.startswith("```"):
            current_num_id = None
            if in_code:
                add_code_block(doc, "\n".join(code_lines), architecture)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith("|"):
            current_num_id = None
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, parse_table(table_lines))
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            current_num_id = None
            level = len(m.group(1))
            add_heading(doc, m.group(2), level, architecture, workflow)
            i += 1
            continue
        if line.startswith("> "):
            current_num_id = None
            add_callout(doc, line[2:].strip())
            i += 1
            continue
        if re.match(r"^-\s+", line):
            current_num_id = None
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^-\s+", "", line))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            if current_num_id is None:
                current_num_id = new_decimal_numbering(doc)
            p = doc.add_paragraph()
            style_paragraph(p, after=4, line=1.25)
            apply_numbering(p, current_num_id)
            add_inline(p, re.sub(r"^\d+\.\s+", "", line))
            i += 1
            continue
        current_num_id = None
        p = doc.add_paragraph()
        style_paragraph(p, after=6, line=1.25)
        add_inline(p, line)
        i += 1

    # Keep the first chapter page break, but avoid an extra page before the first document-control heading.
    for p in doc.paragraphs:
        if p.text == "文档控制":
            p.paragraph_format.page_break_before = False
            break

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
